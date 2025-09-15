import os
import gradio as gr
import requests
import json
from funasr import AutoModel
import torch
from dotenv import load_dotenv
import subprocess
import tempfile
from docx import Document
import re
from fpdf import FPDF

# 加载.env文件中的环境变量
load_dotenv()

# 环境变量配置
GLM45_API_KEY = os.getenv("GLM45_API_KEY", "72c34a2761594dfeb6c9eb501f5c7c11.Wzgr6C7xNI7tqlZq")
GLM45_ENDPOINT = os.getenv("GLM45_ENDPOINT", "https://open.bigmodel.cn/api/paas/v4/chat/completions")
GLM45V_API_KEY = os.getenv("GLM45V_API_KEY", "45897563484447298c2ff4441bb51b34.3mi3ExGuGuPP6i8a")
GLM45V_ENDPOINT = os.getenv("GLM45V_ENDPOINT", "https://open.bigmodel.cn/api/paas/v4/chat/completions")
FUNASR_MODEL_DIR = os.getenv("FUNASR_MODEL_DIR", "iic/speech_paraformer-large-vad-punc_asr_nat-zh-cn-16k-common-vocab8404-pytorch")
USE_GPU = os.getenv("USE_GPU", "True").lower() == "true"

# 初始化FunASR模型
funasr_model = None

def initialize_funasr_model():
    """初始化FunASR模型"""
    global funasr_model
    if funasr_model is None:
        try:
            funasr_model = AutoModel(
                model=FUNASR_MODEL_DIR, 
                vad_model="fsmn-vad", 
                punc_model="ct-punc", 
                device="cuda" if USE_GPU and torch.cuda.is_available() else "cpu"
            )
            # 清理GPU内存
            if torch.cuda.is_available():
                torch.cuda.empty_cache()
                print("FunASR模型初始化后已清理GPU内存缓存")
        except Exception as e:
            print(f"FunASR模型初始化失败: {str(e)}")
            funasr_model = None

def extract_audio_from_video(video_file):
    """从视频中提取音频"""
    if video_file is None:
        print("未提供视频文件")
        return None
    
    try:
        # 创建临时音频文件
        temp_audio_file = tempfile.NamedTemporaryFile(suffix=".wav", delete=False)
        temp_audio_path = temp_audio_file.name
        temp_audio_file.close()
        
        # 使用ffmpeg提取音频，添加-y参数强制覆盖已存在的文件
        cmd = [
            "ffmpeg",
            "-y",  # 添加此参数强制覆盖输出文件
            "-i", video_file,
            "-vn",  # 禁用视频
            "-acodec", "pcm_s16le",  # 音频编解码器
            "-ar", "16000",  # 音频采样率
            "-ac", "1",  # 单声道
            temp_audio_path
        ]
        
        print(f"正在从视频中提取音频: {video_file}")
        result = subprocess.run(cmd, capture_output=True, text=True)
        
        if result.returncode != 0:
            print(f"音频提取失败: {result.stderr}")
            return None
        
        print(f"音频提取完成: {temp_audio_path}")
        return temp_audio_path
    except Exception as e:
        print(f"音频提取出错: {str(e)}")
        return None

def transcribe_audio(audio_file):
    """使用FunASR进行语音识别"""
    if audio_file is None:
        print("未提供音频文件")
        return ""
    
    # 初始化模型（如果尚未初始化）
    print("正在初始化FunASR模型...")
    initialize_funasr_model()
    
    if funasr_model is None:
        print("FunASR模型初始化失败")
        return "FunASR模型未初始化"
    
    try:
        print(f"开始语音识别: {audio_file}")
        res = funasr_model.generate(input=audio_file, batch_size_s=300, hotword=" ")
        transcription = res[0]["text"]
        print(f"语音识别完成，识别结果: {transcription}")
        
        # 清理GPU内存
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
            print("语音识别完成后已清理GPU内存缓存")
        
        return transcription
    except Exception as e:
        error_msg = f"语音识别出错: {str(e)}"
        print(error_msg)
        # 即使出错也尝试清理内存
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
        return error_msg

def analyze_video(video_file, task_context, scoring_prompt=None):
    """使用MiniCPM-V-4_5分析视频内容"""
    print("开始分析视频内容...")
    if video_file is None:
        print("未提供视频文件")
        return "未提供视频文件"
    
    # 导入新的视频分析模块
    try:
        print("正在导入视频分析模块...")
        from minicpm_video_analyzer import analyze_video_with_minicpm
        print("正在调用MiniCPM-V-4_5模型进行视频内容分析...")
        # 将评分标准传递给视频分析函数
        result = analyze_video_with_minicpm(video_file, task_context, scoring_prompt)
        print("视频内容分析完成")
        
        # 清理GPU内存
        torch.cuda.empty_cache()
        print("已清理GPU内存缓存")
        
        return result
    except Exception as e:
        error_msg = f"视频分析出错: {str(e)}"
        print(error_msg)
        # 即使出错也尝试清理内存
        torch.cuda.empty_cache()
        return error_msg

def analyze_audio(audio_file, task_context, scoring_prompt=None):
    """使用GLM-4.5分析音频转录内容"""
    print("开始分析音频内容...")
    if audio_file is None:
        print("未提供音频文件")
        return "未提供音频文件"
    
    # 先进行语音识别
    print("正在进行语音识别...")
    transcription = transcribe_audio(audio_file)
    print(f"语音识别结果: {transcription}")
    
    # 使用传入的评分标准构建提示词，如果没有则使用默认提示词
    if scoring_prompt:
        prompt = f"""请根据以下评分标准分析销售演讲音频转录文本，并结合任务背景提供评分建议。

任务背景：{task_context}

评分标准：{scoring_prompt}

音频转录：{transcription}

请严格按照评分标准进行分析和评分，并提供详细的评分理由和改进建议。"""
    else:
        prompt = f"""请分析以下销售演讲音频转录文本，并结合任务背景提供评分建议。

任务背景：{task_context}

音频转录：{transcription}

请从以下几个维度进行评分（满分10分）：
1. 语言流畅度
2. 销售技巧
3. 表达逻辑

请提供详细的评分理由和改进建议。"""
    
    print("正在调用GLM-4.5模型进行音频内容分析...")
    headers = {
        "Authorization": f"Bearer {GLM45_API_KEY}",
        "Content-Type": "application/json"
    }
    
    data = {
        "model": "glm-4.5",
        "messages": [
            {
                "role": "system",
                "content": "你是一个专业的销售演讲评估专家。"
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        "temperature": 0.6
    }
    
    try:
        response = requests.post(GLM45_ENDPOINT, headers=headers, json=data)
        response.raise_for_status()
        result = response.json()
        answer = result["choices"][0]["message"]["content"]
        print("音频内容分析完成")
        
        # 清理GPU内存（如果使用了GPU）
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
            print("已清理GPU内存缓存")
        
        return answer
    except Exception as e:
        error_msg = f"音频分析出错: {str(e)}"
        print(error_msg)
        # 即使出错也尝试清理内存
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
        return error_msg

def score_presentation(video, task_context, criteria_file=None):
    """主评分函数，支持自定义评分标准"""
    print("开始生成评分报告...")
    print(f"任务背景: {task_context}")
    
    # 读取自定义评分标准
    custom_criteria = None
    if criteria_file is not None:
        custom_criteria = read_scoring_criteria(criteria_file)
    
    # 如果提供了自定义评分标准，使用它；否则使用默认评分标准
    if custom_criteria:
        scoring_prompt = custom_criteria
        print("使用自定义评分标准")
    else:
        # 默认评分标准
        scoring_prompt = f"""请根据以下任务背景对销售人员的视频演讲进行专业评分和分析：
任务背景：{task_context}

请从以下四个维度进行详细分析和评分（每项满分25分）：
1. 面部表情及肢体动作表现（25分）：评估销售人员在视频中的面部表情是否自然、恰当，肢体动作是否得体、有助于表达。
2. 语言表现能力（25分）：评估销售人员的语言表达是否清晰流畅、逻辑性强，语速语调是否恰当。
3. 内容组织与呈现（25分）：评估演讲内容是否条理清晰、重点突出，是否符合销售场景需求。
4. 专业性与说服力（25分）：评估销售人员展现的专业程度和对观众的说服力。

请为每个维度提供具体评分和分析建议，并给出总体评价和改进建议。"""
        print("使用默认评分标准")
    
    # 从视频中提取音频
    audio_file = None
    if video:
        print("正在从视频中提取音频...")
        audio_file = extract_audio_from_video(video)
        if audio_file:
            print("音频提取完成")
        else:
            print("音频提取失败")
    
    print("正在分析视频内容...")
    # 将评分标准传递给视频分析函数
    video_result = analyze_video(video, task_context, scoring_prompt) if video else "未提供视频"
    print("视频分析完成")
    
    print("正在分析音频内容...")
    # 将评分标准传递给音频分析函数
    audio_result = analyze_audio(audio_file, task_context, scoring_prompt) if audio_file else "未提供音频"
    print("音频分析完成")
    
    # 检查音频转录结果是否出错
    if "语音识别出错" in audio_result:
        audio_result = f"{audio_result}，请检查音频文件格式和内容。"
    
    print("正在生成综合评分报告...")
    # 综合评分报告
    full_report = f"""# 销售演讲综合评分报告

## 任务背景
{task_context}

## 综合分析与评分

### 视频表现分析
{video_result}

### 音频表达分析
{audio_result}

## 总体建议
请结合以上各项分析结果，针对性地改进您的销售演讲技巧。

---
*报告生成完成*
"""
    
    print("评分报告生成完成")
    
    # 清理临时音频文件
    if audio_file and os.path.exists(audio_file):
        try:
            os.unlink(audio_file)
            print("临时音频文件已清理")
        except Exception as e:
            print(f"清理临时音频文件失败: {str(e)}")
    
    # 最终清理GPU内存
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
        print("已清理GPU内存缓存")
    
    return full_report

# 修改主函数以使用工作流
def score_presentation_with_download(video, task_context, criteria_file=None):
    """主评分函数，返回报告内容和下载链接，支持自定义评分标准"""
    try:
        # 导入工作流模块
        from workflow import score_presentation_with_workflow
        
        # 使用工作流进行评分
        report_content, report_files = score_presentation_with_workflow(
            video, task_context, criteria_file
        )
        
        return report_content, report_files
    except Exception as e:
        # 如果工作流不可用，回退到原始实现
        print(f"工作流执行出错，回退到原始实现: {str(e)}")
        # 生成报告
        report_content = score_presentation(video, task_context, criteria_file)
        
        # 生成Markdown格式的下载文件
        md_file_path = generate_report_file(report_content)
        
        # 生成PDF格式的下载文件
        pdf_file_path = generate_pdf_report(report_content)
        
        return report_content, [md_file_path, pdf_file_path]

# 添加一个函数来读取DOCX评分标准文件
def read_scoring_criteria(docx_file):
    """读取DOCX格式的评分标准文件"""
    try:
        doc = Document(docx_file)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        print(f"读取评分标准文件时出错: {str(e)}")
        return None

# 添加一个函数来生成PDF报告文件
def generate_pdf_report(report_content, filename=None):
    """生成PDF格式的报告文件供下载"""
    try:
        # 创建临时PDF文件
        if filename is None:
            temp_file = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            filename = temp_file.name
            temp_file.close()
        
        # 将Markdown内容转换为纯文本（简化处理）
        # 移除Markdown标题标记
        text_content = re.sub(r'^#*\s*', '', report_content, flags=re.MULTILINE)
        # 替换Markdown格式
        text_content = re.sub(r'\*\*(.*?)\*\*', r'\1', text_content)  # 粗体
        text_content = re.sub(r'\*(.*?)\*', r'\1', text_content)      # 斜体
        text_content = re.sub(r'`([^`]+)`', r'\1', text_content)      # 行内代码
        
        # 创建简单的PDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)
        
        # 添加标题
        pdf.set_font("Arial", "B", 16)
        pdf.cell(0, 10, "销售演讲综合评分报告", ln=True, align="C")
        pdf.ln(10)
        
        # 添加内容
        pdf.set_font("Arial", size=12)
        # 分行添加内容
        lines = text_content.split('\n')
        for line in lines:
            if line.strip():  # 非空行
                pdf.cell(0, 10, line.strip(), ln=True)
        
        # 保存PDF文件
        pdf.output(filename)
        return filename
    except Exception as e:
        print(f"生成PDF报告时出错: {str(e)}")
        # 出错时创建一个简单的文本文件作为替代
        temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8')
        temp_file.write(report_content)
        temp_file.close()
        return temp_file.name

# 添加一个函数来生成报告文件
def generate_report_file(report_content):
    """生成报告文件供下载"""
    # 创建临时文件
    temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False, encoding='utf-8')
    temp_file.write(report_content)
    temp_file.close()
    return temp_file.name

# Gradio界面
description = """
# AI智能销售演讲评分系统

上传您的销售演讲视频，并提供任务背景，系统将为您生成专业的评分报告。
系统会自动从视频中提取音频进行分析。

您可以选择上传自定义的DOCX格式评分标准文件，系统将根据该标准进行评分。
如果没有上传评分标准文件，将使用默认的评分标准。
"""

iface = gr.Interface(
    fn=score_presentation_with_download,
    inputs=[
        gr.Video(label="演讲视频"),
        gr.Textbox(label="任务背景", lines=3, placeholder="请描述本次销售演讲的任务背景，例如：向IT经理推销企业级软件解决方案"),
        gr.File(label="评分标准文件（可选）", file_types=[".docx"])
    ],
    outputs=[
        gr.Markdown(label="综合评分报告"),
        gr.File(label="下载详细报告（Markdown和PDF格式）", file_count="multiple")
    ],
    title="AI智能销售演讲评分系统",
    description=description,
    examples=[
        [None, "向IT经理推销企业级软件解决方案", None]
    ]
)

if __name__ == "__main__":
    # 初始化FunASR模型
    initialize_funasr_model()
    iface.launch(server_name="0.0.0.0", server_port=7861, show_error=True)