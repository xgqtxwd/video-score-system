"""
使用LangGraph框架改造的销售演讲评分系统工作流
"""
from typing import TypedDict, Optional, List, Dict, Any
from langgraph.graph import StateGraph, END, START
import tempfile
import os
import torch
import requests
from docx import Document
import re
from fpdf import FPDF
from decord import VideoReader, cpu
import numpy as np

# 导入现有的功能模块
from minicpm_video_analyzer import analyze_video_with_minicpm
from app import (
    extract_audio_from_video, 
    transcribe_audio, 
    GLM45_API_KEY, 
    GLM45_ENDPOINT,
    funasr_model,
    initialize_funasr_model
)


class PresentationScoringState(TypedDict):
    """工作流状态定义"""
    video_file: Optional[str]
    audio_file: Optional[str]
    task_context: str
    scoring_criteria: Optional[str]
    video_analysis_result: Optional[str]
    audio_analysis_result: Optional[str]
    final_report: Optional[str]
    report_files: Optional[List[str]]
    error_messages: List[str]


def extract_audio_node(state: PresentationScoringState) -> PresentationScoringState:
    """音频提取节点"""
    print("开始执行音频提取节点...")
    try:
        if state["video_file"]:
            print("正在从视频中提取音频...")
            audio_file = extract_audio_from_video(state["video_file"])
            if audio_file:
                print("音频提取完成")
                state["audio_file"] = audio_file
            else:
                print("音频提取失败")
                state["error_messages"].append("音频提取失败")
        else:
            print("未提供视频文件")
            state["error_messages"].append("未提供视频文件")
    except Exception as e:
        error_msg = f"音频提取出错: {str(e)}"
        print(error_msg)
        state["error_messages"].append(error_msg)
    
    return state


def parse_criteria_node(state: PresentationScoringState) -> PresentationScoringState:
    """评分标准解析节点"""
    print("开始执行评分标准解析节点...")
    # 在实际应用中，这里会从上传的文件中读取评分标准
    # 目前我们保持简化实现
    print("评分标准解析完成")
    return state


def analyze_video_node(state: PresentationScoringState) -> PresentationScoringState:
    """视频分析节点"""
    print("开始执行视频分析节点...")
    try:
        if state["video_file"]:
            print("正在分析视频内容...")
            # 使用传入的评分标准进行分析
            result = analyze_video_with_minicpm(
                state["video_file"], 
                state["task_context"], 
                state["scoring_criteria"]
            )
            state["video_analysis_result"] = result
            print("视频分析完成")
        else:
            state["video_analysis_result"] = "未提供视频文件"
    except Exception as e:
        error_msg = f"视频分析出错: {str(e)}"
        print(error_msg)
        state["error_messages"].append(error_msg)
        state["video_analysis_result"] = error_msg
    
    return state


def analyze_audio_node(state: PresentationScoringState) -> PresentationScoringState:
    """音频分析节点"""
    print("开始执行音频分析节点...")
    try:
        if state["audio_file"]:
            print("正在进行语音识别...")
            transcription = transcribe_audio(state["audio_file"])
            print(f"语音识别结果: {transcription}")
            
            # 使用传入的评分标准构建提示词
            if state["scoring_criteria"]:
                prompt = f"""请根据以下评分标准分析销售演讲音频转录文本，并结合任务背景提供评分建议。

任务背景：{state['task_context']}

评分标准：{state['scoring_criteria']}

音频转录：{transcription}

请严格按照评分标准进行分析和评分，并提供详细的评分理由和改进建议。"""
            else:
                prompt = f"""请分析以下销售演讲音频转录文本，并结合任务背景提供评分建议。

任务背景：{state['task_context']}

音频转录：{transcription}

请从以下几个维度进行评分（满分10分）：
1. 语言流畅度
2. 销售技巧
3. 表达逻辑

请提供详细的评分理由和改进建议。"""
            
            print("正在调用GLM-4.5模型进行音频内容分析...")
            headers = {
                "Authorization": f"Bearer {GLM45_API_KEY}",
                "Content-Type": "application/json"
            }
            
            data = {
                "model": "glm-4.5",
                "messages": [
                    {
                        "role": "system",
                        "content": "你是一个专业的销售演讲评估专家。"
                    },
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],
                "temperature": 0.6
            }
            
            response = requests.post(GLM45_ENDPOINT, headers=headers, json=data)
            response.raise_for_status()
            result = response.json()
            answer = result["choices"][0]["message"]["content"]
            state["audio_analysis_result"] = answer
            print("音频内容分析完成")
        else:
            state["audio_analysis_result"] = "未提供音频文件"
    except Exception as e:
        error_msg = f"音频分析出错: {str(e)}"
        print(error_msg)
        state["error_messages"].append(error_msg)
        state["audio_analysis_result"] = error_msg
    
    return state


def generate_report_node(state: PresentationScoringState) -> PresentationScoringState:
    """报告生成节点"""
    print("开始执行报告生成节点...")
    try:
        # 综合评分报告
        full_report = f"""# 销售演讲综合评分报告

## 任务背景
{state['task_context']}

## 综合分析与评分

### 视频表现分析
{state['video_analysis_result']}

### 音频表达分析
{state['audio_analysis_result']}

## 总体建议
请结合以上各项分析结果，针对性地改进您的销售演讲技巧。

---
*报告生成完成*
"""
        
        state["final_report"] = full_report
        
        # 生成报告文件
        report_files = []
        
        # 生成Markdown格式的下载文件
        temp_file = tempfile.NamedTemporaryFile(mode='w', suffix='.md', delete=False, encoding='utf-8')
        temp_file.write(full_report)
        temp_file.close()
        report_files.append(temp_file.name)
        
        # 生成PDF格式的下载文件
        try:
            # 创建临时PDF文件
            temp_pdf_file = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            filename = temp_pdf_file.name
            temp_pdf_file.close()
            
            # 将Markdown内容转换为纯文本（简化处理）
            # 移除Markdown标题标记
            text_content = re.sub(r'^#*\s*', '', full_report, flags=re.MULTILINE)
            # 替换Markdown格式
            text_content = re.sub(r'\*\*(.*?)\*\*', r'\1', text_content)  # 粗体
            text_content = re.sub(r'\*(.*?)\*', r'\1', text_content)      # 斜体
            text_content = re.sub(r'`([^`]+)`', r'\1', text_content)      # 行内代码
            
            # 创建简单的PDF
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font("Arial", size=12)
            
            # 添加标题
            pdf.set_font("Arial", "B", 16)
            pdf.cell(0, 10, "销售演讲综合评分报告", ln=True, align="C")
            pdf.ln(10)
            
            # 添加内容
            pdf.set_font("Arial", size=12)
            # 分行添加内容
            lines = text_content.split('\n')
            for line in lines:
                if line.strip():  # 非空行
                    pdf.cell(0, 10, line.strip(), ln=True)
            
            # 保存PDF文件
            pdf.output(filename)
            report_files.append(filename)
        except Exception as e:
            print(f"生成PDF报告时出错: {str(e)}")
            # 出错时创建一个简单的文本文件作为替代
            temp_txt_file = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8')
            temp_txt_file.write(full_report)
            temp_txt_file.close()
            report_files.append(temp_txt_file.name)
        
        state["report_files"] = report_files
        print("报告生成完成")
    except Exception as e:
        error_msg = f"报告生成出错: {str(e)}"
        print(error_msg)
        state["error_messages"].append(error_msg)
    
    # 清理临时音频文件
    if state["audio_file"] and os.path.exists(state["audio_file"]):
        try:
            os.unlink(state["audio_file"])
            print("临时音频文件已清理")
        except Exception as e:
            print(f"清理临时音频文件失败: {str(e)}")
    
    # 最终清理GPU内存
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
        print("已清理GPU内存缓存")
    
    return state


def create_workflow():
    """创建工作流"""
    # 创建状态图
    builder = StateGraph(PresentationScoringState)
    
    # 添加节点
    builder.add_node("extract_audio", extract_audio_node)
    builder.add_node("parse_criteria", parse_criteria_node)
    builder.add_node("analyze_video", analyze_video_node)
    builder.add_node("analyze_audio", analyze_audio_node)
    builder.add_node("generate_report", generate_report_node)
    
    # 添加边
    builder.add_edge(START, "extract_audio")
    builder.add_edge("extract_audio", "parse_criteria")
    builder.add_edge("parse_criteria", "analyze_video")
    builder.add_edge("parse_criteria", "analyze_audio")
    builder.add_edge("analyze_video", "generate_report")
    builder.add_edge("analyze_audio", "generate_report")
    builder.add_edge("generate_report", END)
    
    # 编译工作流
    workflow = builder.compile()
    
    return workflow


def score_presentation_with_workflow(video_file: Optional[str], task_context: str, criteria_file: Optional[str] = None):
    """使用工作流进行评分"""
    # 初始化FunASR模型
    initialize_funasr_model()
    
    # 读取评分标准
    scoring_criteria = None
    if criteria_file:
        try:
            doc = Document(criteria_file)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            scoring_criteria = '\n'.join(full_text)
        except Exception as e:
            print(f"读取评分标准文件时出错: {str(e)}")
    
    # 创建初始状态
    initial_state = PresentationScoringState(
        video_file=video_file,
        audio_file=None,
        task_context=task_context,
        scoring_criteria=scoring_criteria,
        video_analysis_result=None,
        audio_analysis_result=None,
        final_report=None,
        report_files=None,
        error_messages=[]
    )
    
    # 创建并执行工作流
    workflow = create_workflow()
    final_state = workflow.invoke(initial_state)
    
    # 返回结果
    return final_state["final_report"], final_state["report_files"]


# 测试代码
if __name__ == "__main__":
    # 这里可以添加测试代码
    pass